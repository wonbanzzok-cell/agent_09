"""Print the structure of a .pptx file: size, layouts, slides, shapes, text, notes.

Also checks the deck against the standard font and, on request, a color theme.

Usage:
    python inspect_pptx.py <file.pptx> [--layouts] [--max-text N] [--font NAME] [--theme pink]

    --layouts     also list every slide layout and its placeholders
    --max-text N  truncate each text to N characters (default 80, 0 = no limit)
    --font NAME   font every run is expected to use (default "Noto Sans KR")
    --theme pink  warn about solid fill and text colors outside the pink theme
"""

import argparse
import sys
from collections import Counter

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.oxml.ns import qn
from pptx.util import Emu

DEFAULT_FONT = "Noto Sans KR"

PINK_THEME = {
    "primary": "E91E63",
    "dark": "AD1457",
    "light": "F8BBD0",
    "background": "FFF0F5",
    "text": "4A2C3A",
    "white": "FFFFFF",
}

THEMES = {"pink": PINK_THEME}

INHERITED = "(inherited)"


class StyleCheck:
    """The font and palette to check against, plus what the deck actually uses."""

    def __init__(self, font, palette):
        self.font = font
        self.palette = palette
        self.fonts = Counter()
        self.colors = Counter()


def inches(value):
    return "?" if value is None else f"{Emu(value).inches:.2f}"


def clip(text, limit):
    text = " / ".join(line for line in text.replace("\v", "\n").splitlines() if line.strip())
    if limit and len(text) > limit:
        return text[:limit] + "..."
    return text


def shape_kind(shape):
    try:
        kind = shape.shape_type
    except NotImplementedError:
        return "UNKNOWN"
    return "UNKNOWN" if kind is None else kind.name


def describe_placeholder(shape):
    fmt = shape.placeholder_format
    return f"placeholder idx={fmt.idx} type={fmt.type.name if fmt.type is not None else '?'}"


def run_fonts(run):
    """Return the Latin and East Asian typefaces set on a run (None = inherited)."""
    latin = run.font.name
    ea = run._r.rPr.find(qn("a:ea"))
    return latin, None if ea is None else ea.get("typeface")


def solid_rgb(fill):
    """Return the hex RGB of a solid fill, or None for any other kind of fill."""
    if fill.type != MSO_FILL.SOLID or fill.fore_color.type != MSO_COLOR_TYPE.RGB:
        return None
    return str(fill.fore_color.rgb)


def check_style(shape, check, detail):
    """Count the fonts and colors a shape uses and warn about the ones off the standard."""
    frames, fills = [], [getattr(shape, "fill", None)]
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        frames.append(shape.text_frame)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                frames.append(cell.text_frame)
                fills.append(cell.fill)

    off_fonts = set()
    for frame in frames:
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                for name in set(run_fonts(run)):
                    check.fonts[name or INHERITED] += 1
                    if name and name != check.font:
                        off_fonts.add(name)
                fills.append(run.font.fill)

    off_colors = set()
    for fill in fills:
        rgb = None if fill is None else solid_rgb(fill)
        if rgb is None:
            continue
        check.colors[rgb] += 1
        if check.palette is not None and rgb not in check.palette.values():
            off_colors.add(rgb)

    if off_fonts:
        print(detail + f"WARNING: font is not {check.font}: {', '.join(sorted(off_fonts))}")
    if off_colors:
        print(detail + f"WARNING: color outside the theme: {', '.join('#' + c for c in sorted(off_colors))}")


def print_shape(shape, limit, slide_w, slide_h, check, indent="    "):
    parts = [f"[{shape.shape_id}] {shape_kind(shape)} \"{shape.name}\""]
    parts.append(
        f"pos=({inches(shape.left)}, {inches(shape.top)}) "
        f"size=({inches(shape.width)} x {inches(shape.height)}) in"
    )
    if shape.is_placeholder:
        parts.append(describe_placeholder(shape))
    print(indent + "  ".join(parts))

    detail = indent + "    "
    if None not in (shape.left, shape.top, shape.width, shape.height) and indent == "    ":
        if (
            shape.left < 0
            or shape.top < 0
            or shape.left + shape.width > slide_w
            or shape.top + shape.height > slide_h
        ):
            print(detail + "WARNING: extends outside the slide")

    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = shape.text_frame.text
        if text.strip():
            print(detail + f"text: {clip(text, limit)}")
        elif shape.is_placeholder:
            print(detail + "WARNING: empty placeholder")

    if getattr(shape, "has_table", False) and shape.has_table:
        table = shape.table
        print(detail + f"table: {len(table.rows)} rows x {len(table.columns)} cols")
        for r, row in enumerate(table.rows):
            cells = [clip(cell.text, 20) for cell in row.cells]
            print(detail + f"  row {r}: {' | '.join(cells)}")

    if getattr(shape, "has_chart", False) and shape.has_chart:
        chart = shape.chart
        try:
            chart_type = chart.chart_type.name
        except Exception:
            chart_type = "UNKNOWN"
        title = chart.chart_title.text_frame.text if chart.has_title else ""
        print(detail + f"chart: {chart_type} title=\"{title}\"")
        for plot in chart.plots:
            categories = [str(c) for c in plot.categories]
            if categories:
                print(detail + f"  categories: {', '.join(categories)}")
            for series in plot.series:
                print(detail + f"  series \"{series.name}\": {list(series.values)}")

    if shape_kind(shape) == "PICTURE":
        try:
            image = shape.image
            print(detail + f"image: {image.content_type} {image.size[0]}x{image.size[1]} px")
        except Exception:
            print(detail + "image: linked or unreadable")

    check_style(shape, check, detail)

    if shape_kind(shape) == "GROUP":
        for child in shape.shapes:
            print_shape(child, limit, slide_w, slide_h, check, indent + "    ")


def print_style_summary(check, theme):
    print(f"\nFonts (expected: {check.font})")
    if not check.fonts:
        print("    (no text runs)")
    for name, count in check.fonts.most_common():
        flag = "" if name in (check.font, INHERITED) else "  WARNING: not the expected font"
        print(f"    {name}: {count} runs{flag}")

    print(f"\nColors (theme: {theme or 'not checked'})")
    if not check.colors:
        print("    (no solid RGB colors)")
    for rgb, count in check.colors.most_common():
        flag = ""
        if check.palette is not None:
            roles = [role for role, value in check.palette.items() if value == rgb]
            flag = f"  {roles[0]}" if roles else "  WARNING: outside the theme"
        print(f"    #{rgb}: {count} uses{flag}")


def main():
    parser = argparse.ArgumentParser(description="Print the structure of a .pptx file.")
    parser.add_argument("path")
    parser.add_argument("--layouts", action="store_true")
    parser.add_argument("--max-text", type=int, default=80)
    parser.add_argument("--font", default=DEFAULT_FONT)
    parser.add_argument("--theme", choices=sorted(THEMES))
    args = parser.parse_args()
    check = StyleCheck(args.font, THEMES.get(args.theme))

    # Windows consoles default to a legacy code page that cannot print Korean.
    sys.stdout.reconfigure(encoding="utf-8")

    prs = Presentation(args.path)
    slide_w, slide_h = prs.slide_width, prs.slide_height

    print(f"File: {args.path}")
    print(f"Slide size: {inches(slide_w)} x {inches(slide_h)} in")
    print(f"Slides: {len(prs.slides)}")
    props = prs.core_properties
    print(f"Title: \"{props.title}\"  Author: \"{props.author}\"")

    if args.layouts:
        print("\nLayouts:")
        for m, master in enumerate(prs.slide_masters):
            for i, layout in enumerate(master.slide_layouts):
                print(f"  master {m} layout {i}: \"{layout.name}\"")
                for ph in layout.placeholders:
                    print(f"      {describe_placeholder(ph)} \"{ph.name}\"")

    for number, slide in enumerate(prs.slides, start=1):
        print(f"\nSlide {number}  (id={slide.slide_id}, layout=\"{slide.slide_layout.name}\")")
        if not len(slide.shapes):
            print("    (no shapes)")
        for shape in slide.shapes:
            print_shape(shape, args.max_text, slide_w, slide_h, check)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                print(f"    notes: {clip(notes, args.max_text)}")

    print_style_summary(check, args.theme)


if __name__ == "__main__":
    main()
