"""Build Anthropic_회사소개_발표자료_v1.pptx. Usage: python build-presentation-v1.py <output.pptx>"""

import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

FONT = "맑은 고딕"
ACCENT = RGBColor(0xE8, 0x6A, 0x10)  # main color: orange
TINT = RGBColor(0xFF, 0xF1, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x26, 0x26, 0x26)
GRAY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xF5, 0xF5, 0xF5)

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
CONTENT_W = SLIDE_W - 2 * MARGIN
BLANK = 6  # index of the blank layout in the default template
TOTAL = 5


def set_font_name(run, name=FONT):
    """Set the Latin, East Asian, and complex-script typefaces of a run.

    font.name writes only <a:latin>, so Korean text would fall back to the
    theme font without the <a:ea> element.
    """
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    anchor = rPr.find(qn("a:latin"))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            anchor.addnext(el)  # schema order is latin, ea, cs
        el.set("typeface", name)
        anchor = el


def style_run(run, size, bold=False, color=DARK):
    set_font_name(run)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def fill_frame(frame, text, size, bold=False, color=DARK, align=PP_ALIGN.LEFT,
               anchor=MSO_ANCHOR.TOP):
    """Write text into a text frame. Pass a list of strings to get one paragraph per item."""
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = align
        run = para.add_run()
        run.text = line
        style_run(run, size, bold, color)


def add_text(slide, left, top, width, height, text, size=20, bold=False,
             color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    fill_frame(box.text_frame, text, size, bold, color, align, anchor)
    return box


def add_rect(slide, left, top, width, height, color, shape=MSO_SHAPE.RECTANGLE):
    rect = slide.shapes.add_shape(shape, left, top, width, height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    rect.line.fill.background()
    rect.shadow.inherit = False
    return rect


def add_slide(prs, number, title, notes):
    """Add a content slide with the title, accent rule, and page number."""
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK])
    add_text(slide, MARGIN, Inches(0.5), CONTENT_W, Inches(0.8), title,
             size=34, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, MARGIN, Inches(1.35), Inches(1.2), Inches(0.08), ACCENT)
    add_text(slide, SLIDE_W - MARGIN - Inches(1.5), Inches(6.75), Inches(1.5), Inches(0.3),
             f"{number} / {TOTAL}", size=12, color=GRAY, align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def build_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = ACCENT
    add_text(slide, MARGIN, Inches(2.3), CONTENT_W, Inches(1.0), "Anthropic",
             size=54, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, MARGIN, Inches(3.45), Inches(1.2), Inches(0.08), WHITE)
    add_text(slide, MARGIN, Inches(3.7), CONTENT_W, Inches(0.7), "Claude를 만드는 AI 회사",
             size=28, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, MARGIN, Inches(6.3), CONTENT_W, Inches(0.5), "회사 소개  |  2026. 10. 10.",
             size=16, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    slide.notes_slide.notes_text_frame.text = (
        "AI 어시스턴트 Claude를 개발한 회사 Anthropic을 소개합니다. "
        "회사 개요, 미션과 핵심 가치, 주요 제품 순서로 말씀드리겠습니다."
    )


def build_overview(prs):
    slide = add_slide(
        prs, 2, "회사 개요",
        "Anthropic은 2021년 OpenAI 출신 연구진이 설립했습니다. 다리오 아모데이가 CEO, "
        "다니엘라 아모데이가 사장을 맡고 있습니다. 주주 이익뿐 아니라 공익을 함께 추구하도록 "
        "정관에 명시한 공익법인(PBC)이며, 장기이익신탁(Long-Term Benefit Trust)이 "
        "이사 선임에 관여하는 지배구조를 갖추고 있습니다.",
    )
    cards = [
        ("설립", "2021년", "다리오·다니엘라 아모데이 남매 등 AI 연구진이 공동 설립"),
        ("본사", "미국 샌프란시스코", "미국·유럽·아시아에 거점을 둔 글로벌 기업"),
        ("기업 형태", "공익법인 (PBC)", "수익과 함께 공익 추구를 정관에 명시"),
        ("사업 영역", "AI 안전 연구·개발", "AI 모델 Claude를 연구하고 제품으로 제공"),
    ]
    card_w, card_h, gap = Inches(5.9), Inches(2.3), Inches(0.333)
    for i, (label, value, desc) in enumerate(cards):
        left = MARGIN + (i % 2) * (card_w + gap)
        top = Inches(1.8) + (i // 2) * (card_h + Inches(0.25))
        add_rect(slide, left, top, card_w, card_h, LIGHT)
        add_rect(slide, left, top, Inches(0.1), card_h, ACCENT)
        text_left, text_w = left + Inches(0.4), card_w - Inches(0.7)
        add_text(slide, text_left, top + Inches(0.2), text_w, Inches(0.4), label,
                 size=16, bold=True, color=ACCENT)
        add_text(slide, text_left, top + Inches(0.65), text_w, Inches(0.7), value,
                 size=28, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, text_left, top + Inches(1.45), text_w, Inches(0.7), desc,
                 size=16, color=GRAY)


def build_mission(prs):
    slide = add_slide(
        prs, 3, "미션과 핵심 가치",
        "Anthropic의 미션은 인류가 변혁적 AI로의 전환을 안전하게 통과하도록 하는 것입니다. "
        "헌법적 AI는 모델이 따를 원칙을 명시해 학습시키는 방식입니다. 책임 있는 확장 정책(RSP)은 "
        "2023년에 발표한 자체 안전 기준으로, 모델의 위험 수준에 따라 요구되는 안전 조치를 정합니다. "
        "해석 가능성 연구는 모델 내부에서 어떤 일이 일어나는지 이해하려는 연구입니다.",
    )
    add_rect(slide, MARGIN, Inches(1.8), CONTENT_W, Inches(1.1), TINT)
    add_text(slide, MARGIN + Inches(0.4), Inches(1.8), CONTENT_W - Inches(0.8), Inches(1.1),
             "미션: 인류가 변혁적 AI로의 전환을 안전하게 통과하도록 한다",
             size=24, bold=True, color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)
    values = [
        ("안전 우선", "성능 경쟁보다 안전성 검증을 앞에 둔다"),
        ("헌법적 AI", "모델이 따를 원칙을 명시하고 그에 맞게 학습시킨다"),
        ("책임 있는 확장 정책 (RSP)", "위험 수준에 따라 안전 조치를 단계적으로 강화한다"),
        ("해석 가능성 연구", "모델 내부의 작동 원리를 이해하는 연구에 투자한다"),
    ]
    for i, (head, desc) in enumerate(values):
        top = Inches(3.2) + i * Inches(0.85)
        add_rect(slide, MARGIN, top + Inches(0.2), Inches(0.25), Inches(0.25), ACCENT)
        add_text(slide, MARGIN + Inches(0.5), top, Inches(4.2), Inches(0.65), head,
                 size=22, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, MARGIN + Inches(4.8), top, CONTENT_W - Inches(4.8), Inches(0.65), desc,
                 size=18, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)


def build_products(prs):
    slide = add_slide(
        prs, 4, "주요 제품",
        "Claude 모델은 성능과 속도에 따라 Opus, Sonnet, Haiku 세 등급으로 나뉩니다. "
        "일반 사용자는 Claude 앱으로, 개발자는 Claude Code와 API로 모델을 사용합니다. "
        "API는 Amazon Bedrock, Google Cloud Vertex AI 같은 클라우드에서도 제공됩니다.",
    )
    rows = [
        ("제품", "설명", "주요 대상"),
        ("Claude 모델", "Opus · Sonnet · Haiku 세 등급의 AI 모델", "모든 제품의 기반"),
        ("Claude 앱", "웹·데스크톱·모바일에서 쓰는 대화형 AI 어시스턴트", "개인, 팀, 기업"),
        ("Claude Code", "터미널과 IDE에서 코드를 읽고 고치는 코딩 에이전트", "개발자"),
        ("Claude API", "서비스에 Claude를 연결하는 개발자 플랫폼", "개발자, 기업"),
    ]
    col_widths = [Inches(2.8), Inches(6.1), CONTENT_W - Inches(8.9)]
    frame = slide.shapes.add_table(len(rows), 3, MARGIN, Inches(1.8), CONTENT_W, Inches(4.5))
    table = frame.table
    for c, width in enumerate(col_widths):
        table.columns[c].width = width
    for r, row in enumerate(rows):
        table.rows[r].height = Inches(0.9)
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.margin_left = cell.margin_right = Inches(0.2)
            if r == 0:
                cell.fill.fore_color.rgb = ACCENT
                fill_frame(cell.text_frame, text, 18, bold=True, color=WHITE,
                           anchor=MSO_ANCHOR.MIDDLE)
            else:
                cell.fill.fore_color.rgb = TINT if r % 2 == 0 else WHITE
                fill_frame(cell.text_frame, text, 18 if c == 0 else 16, bold=(c == 0),
                           color=DARK if c == 0 else GRAY, anchor=MSO_ANCHOR.MIDDLE)


def build_summary(prs):
    slide = add_slide(
        prs, 5, "정리",
        "Anthropic은 안전을 최우선으로 하는 AI 기업입니다. 공익법인이라는 기업 형태와 "
        "안전 연구로 그 원칙을 뒷받침하고, Claude 제품군으로 이를 실제 서비스에 적용하고 있습니다. "
        "감사합니다.",
    )
    points = [
        ("안전을 최우선으로", "AI 안전 연구에서 출발한 회사"),
        ("공익을 지향하는 구조", "공익법인(PBC)으로 운영"),
        ("Claude로 실현", "모델·앱·Claude Code·API 제공"),
    ]
    box_w, box_h, gap = Inches(3.8), Inches(3.2), Inches(0.366)
    for i, (head, desc) in enumerate(points):
        left = MARGIN + i * (box_w + gap)
        top = Inches(1.9)
        add_rect(slide, left, top, box_w, box_h, LIGHT)
        badge = add_rect(slide, left + Inches(0.4), top + Inches(0.4), Inches(0.8), Inches(0.8),
                         ACCENT, MSO_SHAPE.OVAL)
        fill_frame(badge.text_frame, str(i + 1), 24, bold=True, color=WHITE,
                   align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, left + Inches(0.4), top + Inches(1.45), box_w - Inches(0.8), Inches(0.6),
                 head, size=22, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, left + Inches(0.4), top + Inches(2.1), box_w - Inches(0.8), Inches(0.8),
                 desc, size=16, color=GRAY)
    add_text(slide, MARGIN, Inches(5.5), CONTENT_W, Inches(0.8),
             "Anthropic = 안전을 최우선으로 하는 AI 기업",
             size=26, bold=True, color=ACCENT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def build(output_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    prs.core_properties.title = "Anthropic 회사 소개"
    prs.core_properties.subject = "Claude를 만드는 AI 회사"

    build_cover(prs)
    build_overview(prs)
    build_mission(prs)
    build_products(prs)
    build_summary(prs)

    prs.save(output_path)


if __name__ == "__main__":
    build(sys.argv[1])
